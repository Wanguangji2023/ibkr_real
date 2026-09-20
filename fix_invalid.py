# fix_invalid.py
from openpyxl import load_workbook
from config import Config

KEEP = {"ACGN"}   # 只保留真正无效的

wb = load_workbook(Config.INVALID_CONTRACTS)
ws = wb.active

# 从下往上删
for r in range(ws.max_row, 1, -1):
    code = str(ws.cell(row=r, column=1).value or "").strip()
    if code and code not in KEEP:
        ws.delete_rows(r, 1)
        print(f"删除: {code}")

wb.save(Config.INVALID_CONTRACTS)
print("完成")

# 验证
wb = load_workbook(Config.INVALID_CONTRACTS)
ws = wb.active
for r in range(1, ws.max_row + 1):
    print(f"  row{r}:", [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)])