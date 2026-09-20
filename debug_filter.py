# debug_filter.py
import os
from openpyxl import load_workbook
from config import Config

RANGE_ORDER = ["1-2", "2-4", "4-8", "8-16", "16-32", "32-64",
               "64-128", "128-256", "256-512", "512-1024",
               "1024-2048", "2048-4096", "4096-8192", "8192-16384",
               "16384-32768", "32768-65536", "65536-131072",
               "131072-262144", "262144-524288", "524288-1048576",
               "1048576-2097152", "2097152-4194304", "4194304-8388608",
               "8388608-16777216", "16777216-33554432",
               "33554432-67108864", "67108864-134217728",
               "134217728-268435456", "268435456-536870912",
               "536870912-1073741824"]


def _range_ge(rng, min_rng):
    if rng not in RANGE_ORDER or min_rng not in RANGE_ORDER:
        return None   # 未知
    return RANGE_ORDER.index(rng) >= RANGE_ORDER.index(min_rng)


def debug(path):
    print(f"\n{'='*60}\n文件: {path}\n{'='*60}")
    wb = load_workbook(path, data_only=True)

    ws2 = wb[Config.SHEET2_NAME]
    headers2 = [str(ws2.cell(row=1, column=c).value).strip()
                for c in range(1, ws2.max_column + 1)]
    col_buy = headers2.index("双16再分次2") + 1
    print(f"  '双16再分次2' 在第 {col_buy} 列")

    # 1. 收集黄色
    yellow = {}
    for r in range(2, ws2.max_row + 1):
        cell = ws2.cell(row=r, column=1)
        code = cell.value
        if not code:
            continue
        code = str(code).strip()
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            continue
        rgb = str(fill.start_color.rgb).upper() if fill.start_color.rgb else ""
        if len(rgb) == 8:
            rgb = rgb[2:]
        if rgb != "FFFF00":
            continue
        buy_price = ws2.cell(row=r, column=col_buy).value
        yellow[code] = {"buy_price": buy_price, "row": r}

    print(f"\n  黄色行: {len(yellow)} 支")
    for code, info in list(yellow.items())[:20]:
        print(f"    {code:12s} I列值={info['buy_price']!r:20s} (行{info['row']})")

    # 2. Sheet1 C==E + 区间
    ws1 = wb[Config.SHEET1_NAME]
    s1 = {}
    for r in range(2, ws1.max_row + 1):
        code = ws1.cell(row=r, column=1).value
        if not code:
            continue
        code = str(code).strip()
        c_val = str(ws1.cell(row=r, column=3).value or "").strip()
        e_val = str(ws1.cell(row=r, column=5).value or "").strip()
        s1[code] = {"C": c_val, "E": e_val}

    # 3. 逐支分析
    print(f"\n  {'代码':12s} {'C':18s} {'C==E':6s} {'I列':12s} {'≥32-64':8s} {'结论'}")
    for code, info in yellow.items():
        c_val = s1.get(code, {}).get("C", "")
        e_val = s1.get(code, {}).get("E", "")
        ce_match = (c_val == e_val and c_val != "")
        ge = _range_ge(c_val, "32-64")
        buy = info["buy_price"]
        passed = ce_match and ge is True and buy is not None
        status = "✅" if passed else "❌"
        print(f"    {code:12s} {c_val:18s} {str(ce_match):6s} "
              f"{str(buy)[:12]:12s} {str(ge):8s} {status}")


debug(Config.XLSX_ABOVE)
debug(Config.XLSX_BELOW)