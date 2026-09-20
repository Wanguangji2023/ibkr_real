# clean_orgo.py
from openpyxl import load_workbook
from config import Config

wb = load_workbook(Config.EXECUTED_BUYS)
ws = wb.active
for r in range(ws.max_row, 1, -1):
    if str(ws.cell(row=r, column=1).value or "").strip() == "ORGO":
        ws.delete_rows(r, 1)
        print(f"删除 ORGO 行 {r}")
wb.save(Config.EXECUTED_BUYS)
print("完成")