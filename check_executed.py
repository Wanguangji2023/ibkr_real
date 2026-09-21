# check_executed.py
from openpyxl import load_workbook
from config import Config

wb = load_workbook(Config.EXECUTED_BUYS)
ws = wb.active
print("executed_buys.xlsx:")
for r in range(1, ws.max_row + 1):
    print(" ", [ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)])