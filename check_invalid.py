# check_invalid.py
from openpyxl import load_workbook
from config import Config

print("路径:", Config.INVALID_CONTRACTS)
wb = load_workbook(Config.INVALID_CONTRACTS)
ws = wb.active
print(f"行数: {ws.max_row}, 列数: {ws.max_column}")
print(f"表头: {[ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]}")
for r in range(2, ws.max_row + 1):
    print(f"  row{r}: {[ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)]}")