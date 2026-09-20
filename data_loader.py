# data_loader.py
import os
import pandas as pd
from openpyxl import load_workbook
from config import Config

RANGE_ORDER = [
    "1-2", "2-4", "4-8", "8-16", "16-32", "32-64",
    "64-128", "128-256", "256-512", "512-1024",
    "1024-2048", "2048-4096", "4096-8192", "8192-16384",
    "16384-32768", "32768-65536", "65536-131072",
    "131072-262144", "262144-524288", "524288-1048576",
    "1048576-2097152", "2097152-4194304", "4194304-8388608",
    "8388608-16777216", "16777216-33554432",
    "33554432-67108864", "67108864-134217728",
    "134217728-268435456", "268435456-536870912",
    "536870912-1073741824",
]


def _range_ge(rng, min_rng):
    if rng not in RANGE_ORDER or min_rng not in RANGE_ORDER:
        return False
    return RANGE_ORDER.index(rng) >= RANGE_ORDER.index(min_rng)


def load_holding_csv():
    """读取 holding_pnl.csv，只保留美股 + IBKR"""
    path = Config.HOLDING_CSV
    if not os.path.exists(path):
        return {}
    df = pd.read_csv(path, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    result = {}
    for _, row in df.iterrows():
        market = str(row.get("市场", "")).strip()
        platform = str(row.get("下单平台", "")).strip()
        if market != "美股":
            continue
        if platform and platform.upper() != "IBKR":
            continue
        code = str(row.get("代码", "")).strip()
        qty_cost = str(row.get("持有数量@市价", "")).strip()
        cost = row.get("平均成本价", "")
        qty = None
        if qty_cost and "@" in qty_cost:
            try:
                qty = float(qty_cost.split("@")[0].replace(",", ""))
            except Exception:
                qty = None
        try:
            cost = float(str(cost).replace(",", "")) if cost else None
        except Exception:
            cost = None
        result[code] = {"qty": qty, "cost": cost}
    return result


def _compute_i_value(ws, row):
    """
    手动计算 Sheet2 的 I 列（双16再分次2）。
    公式链：
      F = (C+D)/2      双16再分中间
      G = (D+F)/2      双16再分低位中间
      H = (G+D)/2      双16再分次1
      I = (H+D)/2      双16再分次2
    """
    try:
        c = float(ws.cell(row=row, column=3).value)  # C: 最高(双8-16)
        d = float(ws.cell(row=row, column=4).value)  # D: 最低(双8-16)
    except (ValueError, TypeError):
        return None
    f = (c + d) / 2
    g = (d + f) / 2
    h = (g + d) / 2
    i = (h + d) / 2
    return i


def _read_xlsx_pool(path):
    """
    读取一份 xlsx，返回 {ibkr_code: {"range": ..., "buy_price": ...}}
    """
    if not os.path.exists(path):
        return {}

    wb = load_workbook(path, data_only=True)

    # ---- Sheet2 ----
    ws2 = wb[Config.SHEET2_NAME]
    headers2 = [str(ws2.cell(row=1, column=c).value).strip()
                for c in range(1, ws2.max_column + 1)]

    try:
        col_buy = headers2.index("双16再分次2") + 1
    except ValueError:
        col_buy = ord(Config.BUY_PRICE_COLUMN) - ord("A") + 1

    sheet2_rows = {}
    for r in range(2, ws2.max_row + 1):
        cell_a = ws2.cell(row=r, column=1)
        code = cell_a.value
        if not code:
            continue
        code = str(code).strip()

        # 黄色背景
        fill = cell_a.fill
        if fill is None or fill.fill_type != "solid":
            continue
        rgb = fill.start_color.rgb
        if rgb is None:
            continue
        rgb = str(rgb).upper()
        if len(rgb) == 8:
            rgb = rgb[2:]
        if rgb != "FFFF00":
            continue

        # I 列：先读缓存，None 则手动算
        buy_price = ws2.cell(row=r, column=col_buy).value
        if buy_price is None:
            buy_price = _compute_i_value(ws2, r)
        try:
            buy_price = float(buy_price)
        except (ValueError, TypeError):
            continue

        sheet2_rows[code] = {"buy_price": buy_price}

    # ---- Sheet1 ----
    ws1 = wb[Config.SHEET1_NAME]
    sheet1_rows = {}
    for r in range(2, ws1.max_row + 1):
        code = ws1.cell(row=r, column=1).value
        if not code:
            continue
        code = str(code).strip()
        c_val = str(ws1.cell(row=r, column=3).value or "").strip()
        e_val = str(ws1.cell(row=r, column=5).value or "").strip()
        if c_val != e_val:
            continue
        if not _range_ge(c_val, Config.MIN_RANGE):
            continue
        sheet1_rows[code] = {"range": c_val}

    # ---- 合并 ----
    pool = {}
    for code, info in sheet2_rows.items():
        if code in sheet1_rows:
            ibkr_code = code.replace("US.", "")
            pool[ibkr_code] = {
                "xlsx_code": code,
                "range": sheet1_rows[code]["range"],
                "buy_price": info["buy_price"],
            }
    return pool


def load_buy_pool():
    pool = {}
    pool.update(_read_xlsx_pool(Config.XLSX_ABOVE))
    pool.update(_read_xlsx_pool(Config.XLSX_BELOW))
    return pool