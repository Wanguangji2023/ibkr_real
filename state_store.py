import os
from datetime import datetime
from openpyxl import load_workbook, Workbook
from config import Config
import exchange_calendars as xcals

NYSE = xcals.get_calendar("XNYS")

def load_invalid_contracts():
    """返回无效合约代码集合"""
    _ensure_file(Config.INVALID_CONTRACTS,
                 ["代码", "记录时间", "原因"])
    wb = load_workbook(Config.INVALID_CONTRACTS)
    ws = wb.active
    codes = set()
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if code:
            codes.add(str(code).strip())
    return codes


def append_invalid_contract(code, reason=""):
    """写入无效合约黑名单"""
    _ensure_file(Config.INVALID_CONTRACTS,
                 ["代码", "记录时间", "原因"])
    wb = load_workbook(Config.INVALID_CONTRACTS)
    ws = wb.active
    # 去重
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value).strip() == code:
            wb.close()
            return
    ws.append([code, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), reason])
    wb.save(Config.INVALID_CONTRACTS)
# def _ensure_file(path, headers):
#     if not os.path.exists(path):
#         wb = Workbook()
#         ws = wb.active
#         ws.append(headers)
#         wb.save(path)
def _ensure_file(path, headers):
    """确保文件存在且有效；损坏或空文件则重建"""
    need_create = False

    if not os.path.exists(path):
        need_create = True
    elif os.path.getsize(path) < 1000:   # 小于 1KB 视为空/损坏
        need_create = True
        try:
            os.remove(path)
        except Exception:
            pass

    if need_create:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        wb.save(path)
        print(f"[state_store] 已创建: {path} ({os.path.getsize(path)} bytes)")

def _load_codes(path):
    _ensure_file(path, ["代码"])
    wb = load_workbook(path)
    ws = wb.active
    return {str(ws.cell(row=r, column=1).value).strip()
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=1).value}


def load_executed_buys():
    return _load_codes(Config.EXECUTED_BUYS)


def append_executed_buy(code, name, price, qty, amount, status):
    _ensure_file(Config.EXECUTED_BUYS,
                 ["代码", "名称", "买入日期", "买入价", "数量", "金额", "状态"])
    wb = load_workbook(Config.EXECUTED_BUYS)
    ws = wb.active
    ws.append([code, name, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
               price, qty, amount, status])
    wb.save(Config.EXECUTED_BUYS)


def load_silence(active_only=True):
    """返回仍处于静默期的代码集合"""
    _ensure_file(Config.SILENCE_LIST,
                 ["代码", "卖出日期", "静默到期日", "备注"])
    wb = load_workbook(Config.SILENCE_LIST)
    ws = wb.active
    today = datetime.now().date()
    codes = set()
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        expire = ws.cell(row=r, column=3).value
        if not code:
            continue
        code = str(code).strip()
        if not active_only:
            codes.add(code)
            continue
        try:
            expire_date = datetime.strptime(str(expire)[:10], "%Y-%m-%d").date()
            if expire_date >= today:
                codes.add(code)
        except Exception:
            codes.add(code)
    return codes


def append_silence(code, note=""):
    """写入静默期，到期日 = 15 个 NYSE 交易日之后"""
    _ensure_file(Config.SILENCE_LIST,
                 ["代码", "卖出日期", "静默到期日", "备注"])
    wb = load_workbook(Config.SILENCE_LIST)
    ws = wb.active
    today = datetime.now().date()
    session = NYSE.date_to_session(today, direction="next")
    expire_session = NYSE.session_offset(session, Config.SILENCE_DAYS)
    expire_date = expire_session.date()
    ws.append([code, today.strftime("%Y-%m-%d"),
               expire_date.strftime("%Y-%m-%d"), note])
    wb.save(Config.SILENCE_LIST)


def load_loss_watch():
    return _load_codes(Config.LOSS_WATCH)


def append_loss_watch(code, cost, low, max_loss, note=""):
    _ensure_file(Config.LOSS_WATCH,
                 ["代码", "成本价", "最低价", "最大亏损%", "记录日期", "备注"])
    wb = load_workbook(Config.LOSS_WATCH)
    ws = wb.active
    ws.append([code, cost, low, max_loss,
               datetime.now().strftime("%Y-%m-%d %H:%M:%S"), note])
    wb.save(Config.LOSS_WATCH)