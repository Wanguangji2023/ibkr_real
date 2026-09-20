# debug_pool.py
import os
from collections import Counter
from openpyxl import load_workbook
from config import Config


def debug_xlsx(path):
    print(f"\n{'='*60}")
    print(f"文件: {path}")
    print(f"{'='*60}")

    if not os.path.exists(path):
        print("  ❌ 文件不存在")
        return

    print(f"  大小: {os.path.getsize(path)} bytes")

    wb = load_workbook(path, data_only=True)
    print(f"  Sheet 列表: {wb.sheetnames}")

    # ---------- Sheet2 ----------
    if Config.SHEET2_NAME not in wb.sheetnames:
        print(f"  ❌ 找不到 Sheet2: {Config.SHEET2_NAME}")
        return
    ws2 = wb[Config.SHEET2_NAME]
    print(f"\n  --- Sheet2: {Config.SHEET2_NAME} ---")
    print(f"  行数: {ws2.max_row}, 列数: {ws2.max_column}")

    headers = [ws2.cell(row=1, column=c).value
               for c in range(1, ws2.max_column + 1)]
    print(f"  表头前 12 列: {headers[:12]}")

    # 找 双16再分次2
    try:
        idx = headers.index("双16再分次2") + 1
        print(f"  ✅ '双16再分次2' 在第 {idx} 列")
    except ValueError:
        print(f"  ❌ 表头里找不到 '双16再分次2'")
        idx = None

    # A 列颜色分布
    color_counter = Counter()
    yellow_rows = []
    for r in range(2, ws2.max_row + 1):
        cell = ws2.cell(row=r, column=1)
        if not cell.value:
            continue
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            color_counter["no_fill"] += 1
            continue
        rgb = fill.start_color.rgb
        color_counter[str(rgb)] += 1
        # 判断黄色
        rgb_str = str(rgb).upper() if rgb else ""
        if len(rgb_str) == 8:
            rgb_str = rgb_str[2:]
        if rgb_str == "FFFF00":
            yellow_rows.append(cell.value)

    print(f"\n  A 列颜色分布 Top5:")
    for color, count in color_counter.most_common(5):
        print(f"    {color}: {count} 行")
    print(f"  ✅ 黄色行数: {len(yellow_rows)}")
    print(f"     示例: {yellow_rows[:5]}")

    # ---------- Sheet1 ----------
    if Config.SHEET1_NAME not in wb.sheetnames:
        print(f"  ❌ 找不到 Sheet1: {Config.SHEET1_NAME}")
        return
    ws1 = wb[Config.SHEET1_NAME]
    print(f"\n  --- Sheet1: {Config.SHEET1_NAME} ---")
    print(f"  行数: {ws1.max_row}, 列数: {ws1.max_column}")

    headers1 = [ws1.cell(row=1, column=c).value
                for c in range(1, min(ws1.max_column + 1, 8))]
    print(f"  表头前 7 列: {headers1}")

    # C == E 统计
    ce_match = 0
    range_dist = Counter()
    for r in range(2, ws1.max_row + 1):
        c_val = str(ws1.cell(row=r, column=3).value or "").strip()
        e_val = str(ws1.cell(row=r, column=5).value or "").strip()
        if c_val and c_val == e_val:
            ce_match += 1
            range_dist[c_val] += 1
    print(f"  ✅ C==E 行数: {ce_match}")
    print(f"     区间分布: {dict(range_dist)}")


debug_xlsx(Config.XLSX_ABOVE)
debug_xlsx(Config.XLSX_BELOW)