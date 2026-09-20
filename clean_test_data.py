# clean_test_data.py
from openpyxl import load_workbook
from config import Config

# 要删除的测试代码
TEST_CODES = {"TEST1", "TEST2", "TEST3", "TEST4", "TEST"}

for name, path in [
    ("executed_buys", Config.EXECUTED_BUYS),
    ("silence_list", Config.SILENCE_LIST),
    ("loss_watch", Config.LOSS_WATCH),
    ("invalid_contracts", Config.INVALID_CONTRACTS),
]:
    wb = load_workbook(path)
    ws = wb.active
    removed = 0
    # 从下往上删
    for r in range(ws.max_row, 1, -1):
        code = str(ws.cell(row=r, column=1).value or "").strip()
        if code in TEST_CODES:
            ws.delete_rows(r, 1)
            removed += 1
    wb.save(path)
    print(f"{name}: 删除 {removed} 条")

print("\n清理完成")